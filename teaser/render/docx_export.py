"""Optional editable DOCX export of the approved teaser content.

The DOCX carries the same content and disclaimer as the PDF, but Word controls its layout,
so it is not guaranteed to stay on one page after editing. The PDF is the distribution copy.
"""

from __future__ import annotations

import datetime as dt
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from teaser.config import ASSETS
from teaser.render.pdf import ORDER, RenderOptions
from teaser.schema import TeaserDraft


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def render_docx(draft: TeaserDraft, options: RenderOptions | None = None) -> bytes:
    """Build a DOCX with the included content only."""
    options = options or RenderOptions()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.7))
    section.top_margin = section.bottom_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Open Sans"
    normal.font.size = Pt(10)

    primary = _rgb(draft.brand.primary)
    title = doc.add_paragraph()
    run = title.add_run(draft.company_name.value.strip())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(20), primary
    if draft.tagline.include and draft.tagline.value:
        tag = doc.add_paragraph().add_run(draft.tagline.value)
        tag.italic, tag.font.size = True, Pt(11)
    if draft.website.include and draft.website.value:
        doc.add_paragraph().add_run(draft.website.value).font.size = Pt(10)

    for key in ORDER:
        items = [i for i in draft.section_items(key, included_only=True) if i.text.strip()]
        if not items:
            continue
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        h = heading.add_run(draft.headings.get(key, key.value))
        h.bold, h.font.size, h.font.color.rgb = True, Pt(11), primary
        for item in items:
            doc.add_paragraph(item.text, style="List Bullet" if len(items) > 1 else None)

    disclaimer = doc.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(12)
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d = disclaimer.add_run(draft.disclaimer)
    d.italic, d.font.size = True, Pt(9)

    if options.ten_footer:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date: dt.date = options.compiled_on
        text = (f" {draft.company_name.value.strip()} Investor Teaser          1          "
                f"Compiled on {date:%B} {date.day}, {date.year} by TEN Capital Network    ")
        r = footer.add_run(text)
        r.font.name, r.font.size = "Open Sans", Pt(7)
        logo = ASSETS / "TEN_Capital_logo_footer.png"
        if logo.is_file():
            footer.add_run().add_picture(str(logo), width=Inches(0.67), height=Inches(0.25))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
