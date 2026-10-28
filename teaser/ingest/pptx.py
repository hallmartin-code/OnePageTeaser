"""PPTX ingestion with python-pptx: text, tables, groups, image-only slides, theme colours.

Speaker notes are deliberately not read: they are the presenter's private remarks, often
franker than the slides, and were never meant for distribution.
"""

from __future__ import annotations

import io
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape

from teaser.errors import IngestError
from teaser.ingest.models import Deck, Page, first_line, flatten_table, fmt_pages, normalize_text
from teaser.ingest.ocr import ocr_available, ocr_image


def _walk(shape: BaseShape) -> list[BaseShape]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out: list[BaseShape] = []
        for member in shape.shapes:  # type: ignore[attr-defined]
            out.extend(_walk(member))
        return out
    return [shape]


def _theme_colors(data: bytes) -> list[str]:
    """Accent and dark colours from the first theme part, as ``#RRGGBB``."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = sorted(n for n in archive.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n))
            if not names:
                return []
            xml = archive.read(names[0]).decode("utf-8", "replace")
    except (zipfile.BadZipFile, KeyError, OSError):
        return []
    colors: list[str] = []
    for slot in ("accent1", "dk2", "accent2", "accent3"):
        match = re.search(rf"<a:{slot}>\s*<a:srgbClr val=\"([0-9A-Fa-f]{{6}})\"", xml)
        if match:
            colors.append(f"#{match.group(1).upper()}")
    return colors


def load_pptx(data: bytes, filename: str, *, enable_ocr: bool) -> Deck:
    """Read a PPTX from bytes."""
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise IngestError(
            f"{filename} is not a readable PowerPoint file ({exc}). If it opens in "
            f"PowerPoint, re-save it as .pptx or export it to PDF."
        ) from exc

    can_ocr = enable_ocr and ocr_available()
    pages: list[Page] = []
    logos: list[bytes] = []
    ocr_pages: list[int] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        pictures: list[bytes] = []
        has_table = False
        for top in slide.shapes:
            for shape in _walk(top):
                if getattr(shape, "has_table", False):
                    has_table = True
                    rows = [[c.text for c in row.cells] for row in shape.table.rows]  # type: ignore[attr-defined]
                    parts.append("[table]")
                    parts.extend(flatten_table(rows))
                elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():  # type: ignore[attr-defined]
                    parts.append(shape.text_frame.text)  # type: ignore[attr-defined]
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        pictures.append(shape.image.blob)  # type: ignore[attr-defined]
                    except Exception:
                        pass
        text = normalize_text("\n".join(parts))
        source = "text" if text else "none"
        if len(text) < 20 and pictures and can_ocr:
            recognised = normalize_text("\n".join(ocr_image(b) for b in pictures))
            if recognised:
                text = normalize_text(f"{text}\n{recognised}")
                source = "ocr" if source == "none" else "mixed"
                ocr_pages.append(number)
        if number == 1:
            logos = [b for b in pictures[:4] if len(b) < 3_000_000]

        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip()[:200] if title_shape is not None else ""
        pages.append(
            Page(
                index=number,
                title=title or first_line(text),
                text=text,
                text_source=source,  # type: ignore[arg-type]
                image_count=len(pictures),
                has_table=has_table,
            )
        )

    if not pages:
        raise IngestError(f"{filename} contains no slides.")

    deck = Deck(
        filename=filename,
        source_format="pptx",
        pages=pages,
        theme_colors=_theme_colors(data),
        logo_candidates=logos,
    )
    if ocr_pages:
        deck.warnings.append(
            f"Text on slide(s) {fmt_pages(ocr_pages)} was recovered by OCR from images and may "
            f"contain recognition errors."
        )
    unreadable = deck.unreadable_pages
    if unreadable:
        hint = (
            "OCR found no legible text in their images."
            if can_ocr
            else "Enable OCR (and install rapidocr-onnxruntime) to try reading their images."
        )
        deck.warnings.append(
            f"Slide(s) {fmt_pages(unreadable)} have little or no readable text "
            f"(images, charts, or diagrams). {hint} Charts embedded as pictures cannot be read."
        )
    return deck
